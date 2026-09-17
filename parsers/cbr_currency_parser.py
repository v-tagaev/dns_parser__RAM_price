import requests
import pandas as pd
import json
import os
from datetime import datetime, timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from save_parsers_to_db import get_engine   #общее подключение к БД postgres и парсер
#логгирование
import logging
from setup_logger import setup_logger
logger = setup_logger(__name__)


def main():
    url = "https://www.cbr-xml-daily.ru/daily_json.js"
    resp = requests.get(url)
    resp.raise_for_status()
    logger.info(f'Поключаемся к {url}, ответ сайта: {resp.status_code}\n')
    data = resp.json()

    valute_rows1 = []
    valute_rows2 = []

    rate_date = datetime.fromisoformat(data['Timestamp']).replace(tzinfo=None)
    for i in data['Valute'].values():        
        charcode = i['CharCode']
        numcode = str(i['NumCode'])
        nominal = int(i['Nominal'])
        currency_name = i['Name']
        rate = float(i['Value']/nominal)
    
        valute_rows1.append({
            'currency_numcode': numcode,     #код валюты ЦБ
            'currency_charcode': charcode,   #код валюты буквенный типо USD
            'currency_name': currency_name,   #название валюты
        })
        
        valute_rows2.append({
                'currency_charcode': charcode,   #код валюты буквенный типо USD
                'rate_nominal': nominal,     #номинал валюты
                'rate': rate,           #курс валюты
                'rate_date': rate_date
            })
    #print(f'{valute_rows1}\n')
    #print(f'{valute_rows2}\n')
    df1 = pd.DataFrame(valute_rows1) #-- pd-Pandas DataFrame-двумерная таблица    
    df2 = pd.DataFrame(valute_rows2) #-- pd-Pandas DataFrame-двумерная таблица
    #------------------------------------------------------------------------
    #Проверяем работаем в Docker или Windows сохраняем либо в БД, либо в Excel (windows)
    #------------------------------------------------------------------------
    is_docker = "AIRFLOW_HOME" in os.environ

    if is_docker:
        logger.info(" Парсер был запущен в Docker/Airflow. Сохраняем в БД")

        engine = get_engine() #подключаемся к БД
        
        with engine.begin() as connection:
        # 1. Вставляем валюты, игнорируя уже существующие
            for row in valute_rows1:
                connection.execute(
                    text("""
                        INSERT INTO currency (currency_numcode, currency_charcode, currency_name)
                        VALUES (:numcode, :charcode, :name)
                        ON CONFLICT DO NOTHING
                    """),
                    {
                        "numcode": row['currency_numcode'],
                        "charcode": row['currency_charcode'],
                        "name": row['currency_name'],
                    }
                )

            # 2. Получаем mapping charcode -> currency_id
            mapping = pd.read_sql(
                "SELECT currency_id, currency_charcode FROM currency",
                connection
            )
            code_to_id = dict(zip(mapping['currency_charcode'], mapping['currency_id']))

            # 3. Добавляем currency_id в df2
            df2['currency_id'] = df2['currency_charcode'].map(code_to_id)

            # 4. Вставляем курсы
            df2_to_insert = df2[['currency_id', 'rate_nominal', 'rate', 'rate_date']]
            df2_to_insert.to_sql('exchange_rate', connection, if_exists='append', index=False)

            logger.info(f'Сохранено {len(df2_to_insert)} строк курсов')
    else:
        logger.info(" Парсер был запущен на Windows. Сохраняем в Excel")
        #--название для excel файла
        output_file = f'valute_rows_{rate_date.strftime("%Y.%m.%d_%H-%M")}.xlsx'

        if not df2.empty:
                                
                df_excel = df2.copy()
                df_excel.to_excel(output_file, index=False)
                
                wb = load_workbook(output_file)
                ws = wb.active
                for col_idx, col in enumerate(ws.columns, 1):
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col_idx)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
                wb.save(output_file)
                logger.info(f"Данные успешно сохранены в файл: {output_file}")
        else:
            logger.error("!! Не удалось собрать данные.")

if __name__ == '__main__':
    main()